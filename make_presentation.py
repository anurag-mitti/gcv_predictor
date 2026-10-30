import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # 100% High-Contrast Colors - Absolutely NO black or dark gray text
    BG_COLOR = RGBColor(15, 23, 42)         # Deep slate navy (#0F172A)
    CARD_BG = RGBColor(30, 41, 59)          # Card background (#1E293B)
    TEXT_WHITE = RGBColor(255, 255, 255)    # 100% Pure Crisp White
    ACCENT_CYAN = RGBColor(56, 189, 248)    # Sky Cyan (#38BDF8)
    ACCENT_GREEN = RGBColor(52, 211, 153)   # Emerald Green (#34D399)
    ACCENT_AMBER = RGBColor(251, 146, 60)   # Vivid Amber (#FB923C)
    
    blank_layout = prs.slide_layouts[6]
    
    def apply_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="RESEARCH & ENGINEERING WALKTHROUGH"):
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        p_cat = tb_cat.text_frame.paragraphs[0]
        run_cat = p_cat.add_run()
        run_cat.text = category.upper()
        run_cat.font.size = Pt(11)
        run_cat.font.bold = True
        run_cat.font.color.rgb = ACCENT_CYAN

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        p = tb.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = title_text
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = TEXT_WHITE

    def add_run_line(paragraph, text, size=12, color=TEXT_WHITE, bold=False, space_before=0):
        if space_before > 0:
            paragraph.space_before = Pt(space_before)
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return run

    # ==========================================
    # SLIDE 1: TITLE SLIDE (ALL TEXT BRIGHT WHITE & CYAN)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_bg(s1)

    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.18), Inches(3.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_CYAN
    bar.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.2), Inches(4.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    add_run_line(p1, "COAL GROSS CALORIFIC VALUE (GCV) PREDICTION", size=36, color=TEXT_WHITE, bold=True)

    p2 = tf1.add_paragraph()
    add_run_line(p2, "High-Precision Estimation from Proximate Analysis via Bayesian-Tuned Stacking Ensembles", size=20, color=ACCENT_CYAN, bold=True, space_before=12)

    p3 = tf1.add_paragraph()
    bullets_s1 = [
        "• Optimized Machine Learning Architecture (XGBoost, Random Forest, SVR + Ridge Meta-Learner)",
        "• Zero-Leakage Data Hygiene & ASTM D388 Coal Rank Specialization",
        "• Model Interpretability via SHAP (SHapley Additive exPlanations)",
        "• Live Interactive Web Deployment with Chemical Safety Guardrails"
    ]
    p3.space_before = Pt(20)
    for b in bullets_s1:
        run_b = p3.add_run()
        run_b.text = b + "\n"
        run_b.font.size = Pt(14)
        run_b.font.bold = False
        run_b.font.color.rgb = TEXT_WHITE  # Changed from black/muted to pure crisp white!

    # ==========================================
    # SLIDE 2: PROBLEM STATEMENT & MOTIVATION
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_bg(s2)
    add_header(s2, "Industrial Motivation & Problem Formulation", "CONTEXT & OBJECTIVE")

    cards = [
        ("The Laboratory Bottleneck", 
         "Standard bomb calorimetry requires oxygen-pressurized combustion bombs, calibrated water jackets, and skilled technicians.\n\nResult: 24–48 hour turnaround times delay critical coal blending decisions in power plants.",
         ACCENT_AMBER),
        ("The Proximate Advantage",
         "Proximate analysis (Moisture, Volatile Matter, Ash) is rapid, inexpensive, and standardized.\n\nGoal: Map proximate parameters directly to GCV with high fidelity, enabling immediate day-of-delivery decision making.",
         ACCENT_CYAN),
        ("Rigorous Target Metrics",
         "Target RMSE: <= 200 Kcal/kg\nTarget CV Score: < 8.0%\n\nOur system surpassed both tolerances, delivering 150.28 Kcal/kg (2.18% CV) overall and 134.18 Kcal/kg (1.91% CV) for bituminous coal.",
         ACCENT_GREEN)
    ]

    for i, (head, body, acc) in enumerate(cards):
        left = Inches(0.8 + i * 3.9)
        top = Inches(1.8)
        w = Inches(3.7)
        h = Inches(4.8)
        
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = acc
        card.line.width = Pt(1.5)

        tb = s2.shapes.add_textbox(left + Inches(0.25), top + Inches(0.25), w - Inches(0.5), h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p_h = tf.paragraphs[0]
        add_run_line(p_h, head, size=18, color=acc, bold=True)

        p_b = tf.add_paragraph()
        add_run_line(p_b, body, size=13, color=TEXT_WHITE, bold=False, space_before=14)

    # ==========================================
    # SLIDE 3: DATA HYGIENE & PREPROCESSING
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_bg(s3)
    add_header(s3, "Data Hygiene & Zero-Target-Leakage Protocol", "DATASET & PREPROCESSING")

    steps = [
        ("1. Unit Standardization", "Raw dataset contains imperial heating values in Btu/lb. Converted to international metric standard:\nGCV (Kcal/kg) = Btu / 1.8", ACCENT_CYAN),
        ("2. Zero Target Leakage", "In amateur ML workflows, imputers often fit across features AND target GCV. We enforced strict isolation: KNN imputation fits ONLY on features [Moisture, VM, Ash].", ACCENT_GREEN),
        ("3. Physical Bounds", "Enforced physical combustion realism:\n• GCV: 2000 to 9000 Kcal/kg\n• Standard Ash <= 55%\n• Moisture <= 45%, VM <= 60%", ACCENT_AMBER),
        ("4. Mass Balance Integrity", "By proximate definition, components sum to 100%:\nFixed Carbon = 100 - (Moisture + VM + Ash)\nRetained exclusively samples with Fixed Carbon > 0.", ACCENT_CYAN)
    ]

    for i, (title, desc, color) in enumerate(steps):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.8 + row * 2.5)
        w = Inches(5.6)
        h = Inches(2.2)

        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = color
        box.line.width = Pt(1.5)

        tb = s3.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), w - Inches(0.5), h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        add_run_line(p_t, title, size=16, color=color, bold=True)

        p_d = tf.add_paragraph()
        add_run_line(p_d, desc, size=12.5, color=TEXT_WHITE, bold=False, space_before=8)

    # ==========================================
    # SLIDE 4: FEATURE ENGINEERING (FUEL RATIO)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_bg(s4)
    add_header(s4, "Domain Feature Engineering: Why Add Fuel Ratio?", "FEATURE ENGINEERING")

    tb_left = s4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.8))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    add_run_line(p, "The Mathematical Limitation of Tree Models:", size=18, color=ACCENT_AMBER, bold=True)

    points = [
        "Decision trees split orthogonally (Xi > c). They cannot naturally compute division between features without hundreds of step-wise splits.",
        "In coal petrology, the Fuel Ratio is defined as:\nFuel Ratio = Fixed Carbon / Volatile Matter",
        "Fuel Ratio is the universally accepted chemical proxy for coal rank (Anthracite vs Bituminous vs Lignite).",
        "By calculating it explicitly prior to training, we feed the model an explicit rank descriptor without requiring extra lab testing."
    ]
    for pt in points:
        p = tf_l.add_paragraph()
        add_run_line(p, "• " + pt, size=13, color=TEXT_WHITE, bold=False, space_before=10)

    # Formula Card Right
    fcard = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(1.8), Inches(5.3), Inches(4.8))
    fcard.fill.solid()
    fcard.fill.fore_color.rgb = CARD_BG
    fcard.line.color.rgb = ACCENT_GREEN
    fcard.line.width = Pt(1.5)

    tb_r = s4.shapes.add_textbox(Inches(7.5), Inches(2.1), Inches(4.7), Inches(4.2))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_r1 = tf_r.paragraphs[0]
    add_run_line(p_r1, "Mass Balance Formulation:", size=17, color=ACCENT_GREEN, bold=True)

    p_r2 = tf_r.add_paragraph()
    add_run_line(p_r2, "FC = 100 - (Moisture + VM + Ash)\n\nFuel Ratio = FC / VM", size=16, color=TEXT_WHITE, bold=True, space_before=12)

    p_r3 = tf_r.add_paragraph()
    add_run_line(p_r3, "\nImpact on Training:\n• Drastically simplifies tree depth requirements.\n• Enables shallow XGBoost trees (depth 4) to capture coal rank variations.\n• Consistently reduces CV error across all folds.", size=13, color=TEXT_WHITE, bold=False, space_before=12)

    # ==========================================
    # SLIDE 5: BASELINE BENCHMARKING
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_bg(s5)
    add_header(s5, "Quantitative Benchmarking: Proving Ensemble Superiority", "MODEL EVALUATION")

    rows, cols = 6, 4
    left, top, width, height = Inches(0.8), Inches(1.8), Inches(6.5), Inches(4.5)
    table_shape = s5.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.6)
    table.columns[1].width = Inches(1.3)
    table.columns[2].width = Inches(1.3)
    table.columns[3].width = Inches(1.3)

    headers = ["Model Architecture", "5-Fold RMSE", "5-Fold MAE", "R² Score"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 58, 138)
        p = cell.text_frame.paragraphs[0]
        add_run_line(p, h, size=11, color=TEXT_WHITE, bold=True)

    data = [
        ("Support Vector Machine (SVR)", "612.45", "487.12", "0.3541"),
        ("Decision Tree Regressor", "217.30", "148.80", "0.9168"),
        ("Random Forest Regressor", "154.84", "109.12", "0.9571"),
        ("XGBoost Regressor", "152.14", "107.45", "0.9585"),
        ("Stacked Ensemble (Ours)", "150.28", "105.97", "0.9595")
    ]

    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.fill.solid()
            if i == 4:
                cell.fill.fore_color.rgb = RGBColor(6, 78, 59)
            else:
                cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            add_run_line(p, val, size=11, color=TEXT_WHITE, bold=(i == 4))

    comp_plot = "baselines/model_comparison_plot.png"
    if os.path.exists(comp_plot):
        s5.shapes.add_picture(comp_plot, Inches(7.6), Inches(1.8), width=Inches(5.0))

    # ==========================================
    # SLIDE 6: ENSEMBLE ARCHITECTURE WITH DIAGRAM PIC
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_bg(s6)
    add_header(s6, "End-to-End System & Stacking Architecture", "SYSTEM ARCHITECTURE")

    arch_img = "models/architecture_diagram.png"
    if os.path.exists(arch_img):
        # Insert high-resolution visual architecture diagram
        s6.shapes.add_picture(arch_img, Inches(0.8), Inches(1.6), width=Inches(11.73))

    # ==========================================
    # SLIDE 7: BAYESIAN OPTIMIZATION (OPTUNA)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_bg(s7)
    add_header(s7, "Bayesian Hyperparameter Optimization via Optuna", "HYPERPARAMETER TUNING")

    tb_left7 = s7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.8), Inches(4.8))
    tf7 = tb_left7.text_frame
    tf7.word_wrap = True

    p = tf7.paragraphs[0]
    add_run_line(p, "Why Bayesian Optimization over Grid Search?", size=18, color=ACCENT_CYAN, bold=True)

    points7 = [
        "Grid Search is computationally expensive and samples uninformative parameter regions.",
        "Optuna utilizes Tree-structured Parzen Estimators (TPE) to model the probability distribution of 5-fold CV RMSE.",
        "It concentrates evaluation trials where parameter combinations historically yield lower error.",
        "Optimal parameters saved to models/best_params.json for transparent, reproducible deployment."
    ]
    for pt in points7:
        p = tf7.add_paragraph()
        add_run_line(p, "• " + pt, size=13, color=TEXT_WHITE, bold=False, space_before=10)

    # Param Box Right
    pbox = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8))
    pbox.fill.solid()
    pbox.fill.fore_color.rgb = CARD_BG
    pbox.line.color.rgb = ACCENT_GREEN
    pbox.line.width = Pt(1.5)

    tb_p = s7.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True

    p = tf_p.paragraphs[0]
    add_run_line(p, "Tuned Hyperparameter Snapshot:", size=16, color=ACCENT_GREEN, bold=True)

    params_text = (
        "XGBoost Regressor:\n"
        "• n_estimators = 300  |  max_depth = 4\n"
        "• learning_rate = 0.0256  |  subsample = 0.7\n"
        "• colsample_bytree = 0.8  |  min_child_weight = 2\n\n"
        "Random Forest Regressor:\n"
        "• n_estimators = 200  |  max_depth = 15\n"
        "• min_samples_leaf = 4  |  max_features = 0.7\n\n"
        "Support Vector Regressor (SVR):\n"
        "• C = 1.292  |  epsilon = 0.0857  |  gamma = 'auto'\n\n"
        "Stacking Meta-Learner (Ridge):\n"
        "• alpha = 0.0105  (Optimal L2 shrinkage)"
    )
    p2 = tf_p.add_paragraph()
    add_run_line(p2, params_text, size=11.5, color=TEXT_WHITE, bold=False, space_before=10)

    # ==========================================
    # SLIDE 8: BITUMINOUS SPECIALIZATION
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_bg(s8)
    add_header(s8, "Coal Rank Specialization: ASTM D388 Bituminous Model", "RANK SPECIALIZATION")

    tb_b_left = s8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.8), Inches(4.8))
    tf_b = tb_b_left.text_frame
    tf_b.word_wrap = True

    p = tf_b.paragraphs[0]
    add_run_line(p, "The Scientific Rationale for Rank Segregation:", size=17, color=ACCENT_GREEN, bold=True)

    b_pts = [
        "Coal is chemically non-homogeneous across geological ranks. Lignite, Bituminous, and Anthracite exhibit vastly different volatile release kinetics.",
        "Under ASTM D388, Bituminous Coal is strictly classified by:\n• Moist Mineral-Matter-Free Btu >= 10,500 Btu/lb\n• Dry MMF Volatile Matter >= 14.0%\n• Dry MMF Fixed Carbon < 86.0%",
        "We isolated 4,494 verified bituminous samples and trained a dedicated model pipeline (src/train_bituminous.py).",
        "Performance Impact:\n• RMSE dropped to 134.18 Kcal/kg (vs 150.28)\n• MAE dropped to 96.29 Kcal/kg (sub-100 error!)\n• CV Score reached 1.91%"
    ]
    for pt in b_pts:
        p = tf_b.add_paragraph()
        add_run_line(p, "• " + pt, size=12.5, color=TEXT_WHITE, bold=False, space_before=8)

    diag_bit = "models/actual_vs_predicted_bituminous.png"
    if os.path.exists(diag_bit):
        s8.shapes.add_picture(diag_bit, Inches(7.0), Inches(1.8), width=Inches(5.5))

    # ==========================================
    # SLIDE 9: EXPLAINABLE AI (SHAP)
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_bg(s9)
    add_header(s9, "Explainable AI: SHAP Chemical Interpretability", "MODEL INTERPRETABILITY")

    tb_s_left = s9.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.8), Inches(4.8))
    tf_s = tb_s_left.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    add_run_line(p, "Overcoming the 'Black Box' Critique:", size=17, color=ACCENT_CYAN, bold=True)

    shap_pts = [
        "Engineers rightfully distrust black-box ML models when plant safety and boiler efficiency are at stake.",
        "We integrated SHAP (SHapley Additive exPlanations) using TreeExplainer to explain individual local predictions.",
        "Thermodynamic Sanity Check:\n• Moisture: Strong negative SHAP driver (water absorbs heat via vaporization at 540 cal/g).\n• Ash: Strong negative SHAP driver (inert mineral matter acts as a thermal heat sink).\n• Fixed Carbon & Fuel Ratio: Dominant positive drivers (primary energy-producing organic matrix).",
        "The model's internal feature attribution perfectly validates established combustion thermodynamics."
    ]
    for pt in shap_pts:
        p = tf_s.add_paragraph()
        add_run_line(p, "• " + pt, size=12, color=TEXT_WHITE, bold=False, space_before=8)

    shap_img = "models/shap_waterfall.png"
    if os.path.exists(shap_img):
        s9.shapes.add_picture(shap_img, Inches(7.0), Inches(1.8), width=Inches(5.5))

    # ==========================================
    # SLIDE 10: WEB DASHBOARD & PRODUCTION DEPLOYMENT
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_bg(s10)
    add_header(s10, "Interactive Streamlit Deployment & Domain Guardrails", "SYSTEM DEPLOYMENT")

    dep_cards = [
        ("Dual-Tab Architecture", "• Tab 1: Universal Model (All coal ranks)\n• Tab 2: Bituminous Specialized Model (ASTM D388)\nUsers seamlessly switch between general coal and high-precision bituminous mode.", ACCENT_CYAN),
        ("Domain-Specific Guardrails", "• Mass Balance Validation: Flags an immediate error if Moisture + VM + Ash > 100%.\n• Ash Limit Rule: Rejects inputs exceeding realistic combustible limits (Ash > 55%).", ACCENT_AMBER),
        ("Real-Time Explainability", "Every slider adjustment instantly triggers model re-prediction and renders dynamic SHAP waterfall plots explaining the chemical drivers behind that specific prediction.", ACCENT_GREEN),
        ("Production Cloud Ready", "Deployed via Streamlit Community Cloud with GitHub CI/CD tracking main branch. Serialized models (<20MB) load in sub-second inference speed.", ACCENT_CYAN)
    ]

    for i, (title, body, color) in enumerate(dep_cards):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.8 + row * 2.5)
        w = Inches(5.6)
        h = Inches(2.2)

        box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = color
        box.line.width = Pt(1.5)

        tb = s10.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), w - Inches(0.5), h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        add_run_line(p_t, title, size=16, color=color, bold=True)

        p_d = tf.add_paragraph()
        add_run_line(p_d, body, size=12, color=TEXT_WHITE, bold=False, space_before=8)

    # ==========================================
    # SLIDE 11: SUMMARY & FUTURE SCOPE
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    apply_bg(s11)
    add_header(s11, "Project Summary & Future Research Directions", "CONCLUSION & ROADMAP")

    sbox = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    sbox.fill.solid()
    sbox.fill.fore_color.rgb = CARD_BG
    sbox.line.color.rgb = ACCENT_GREEN
    sbox.line.width = Pt(1.5)

    tb_s = s11.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    add_run_line(p, "Final Performance Summary:", size=17, color=ACCENT_GREEN, bold=True)

    summary_text = (
        "Universal Stacking Model (All Ranks):\n"
        "• 5-Fold CV RMSE: 150.28 Kcal/kg\n"
        "• 5-Fold CV MAE: 105.97 Kcal/kg\n"
        "• Coefficient of Variation (CV): 2.18%\n"
        "• Variance Explained (R²): 0.9595\n\n"
        "Bituminous Specialized Model (ASTM D388):\n"
        "• 5-Fold CV RMSE: 134.18 Kcal/kg  (-16.10)\n"
        "• 5-Fold CV MAE: 96.29 Kcal/kg  (<100!)\n"
        "• Coefficient of Variation (CV): 1.91%\n"
        "• Variance Explained (R²): 0.9524"
    )
    p2 = tf_s.add_paragraph()
    add_run_line(p2, summary_text, size=12, color=TEXT_WHITE, bold=False, space_before=10)

    fbox = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    fbox.fill.solid()
    fbox.fill.fore_color.rgb = CARD_BG
    fbox.line.color.rgb = ACCENT_CYAN
    fbox.line.width = Pt(1.5)

    tb_f = s11.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True

    p = tf_f.paragraphs[0]
    add_run_line(p, "Future Research & Engineering Scope:", size=17, color=ACCENT_CYAN, bold=True)

    future_text = (
        "1. Dual-Tier Elemental Analysis Mode:\n"
        "Ingest Ultimate Analysis (Carbon, Hydrogen, Oxygen, Sulfur) when available to drive RMSE below 50 Kcal/kg.\n\n"
        "2. Quantile Regression & Uncertainty Bounds:\n"
        "Predict 90% confidence intervals around GCV values rather than point estimates to inform risk-aware procurement.\n\n"
        "3. High-Throughput Batch Processing:\n"
        "Enable bulk CSV uploads for trainload shipments with automated PDF laboratory certificates.\n\n"
        "4. Classical Formula Benchmarking:\n"
        "Provide direct mathematical comparisons against Goutal and Dulong empirical equations."
    )
    p2 = tf_f.add_paragraph()
    add_run_line(p2, future_text, size=12, color=TEXT_WHITE, bold=False, space_before=10)

    out_filename = "Coal_GCV_Prediction_Presentation.pptx"
    prs.save(out_filename)
    print(f"Presentation successfully regenerated: {out_filename}")

if __name__ == '__main__':
    create_deck()
